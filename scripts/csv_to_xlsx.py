import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
import os
import string

# Set input and output directories
input_dir = "."
output_dir = "./xlsx"
os.makedirs(output_dir, exist_ok=True)

# Loop through all CSV files in the input directory
for filename in os.listdir(input_dir):
    if filename.endswith(".csv"):
        csv_path = os.path.join(input_dir, filename)
        xlsx_path = os.path.join(output_dir, filename.replace(".csv", ".xlsx"))
        
        try:
            # Read CSV (adjust delimiter if needed)
            df = pd.read_csv(csv_path, sep=";", engine="python")

            # Write to Excel without index
            df.to_excel(xlsx_path, index=False)

            # Load workbook and worksheet
            wb = load_workbook(xlsx_path)
            ws = wb.active

            # Define table range
            num_cols = len(df.columns)
            end_col_letter = string.ascii_uppercase[num_cols - 1]
            end_row = df.shape[0] + 1  # header + data
            table_range = f"A1:{end_col_letter}{end_row}"

            # Create table with style
            table = Table(displayName="Table_" + filename.replace(".csv", ""), ref=table_range)
            style = TableStyleInfo(
                name="TableStyleMedium9",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False
            )
            table.tableStyleInfo = style
            ws.add_table(table)

            # Save formatted Excel file
            wb.save(xlsx_path)
            print(f"[✓] Converted and formatted: {filename} → {xlsx_path}")

        except Exception as e:
            print(f"[!] Failed to convert {filename}: {e}")
